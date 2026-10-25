"""Build the voice-over spreadsheet from a phrase list exported from the game (phrases.json)."""
import json,sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
src,dst=sys.argv[1],sys.argv[2]
rows=json.load(open(src))
seen={}
for sec,t,k in rows:
    if t not in seen: seen[t]={'secs':[sec],'kind':k}
    else:
        if sec not in seen[t]['secs']: seen[t]['secs'].append(sec)
        if not seen[t]['kind'] and k: seen[t]['kind']=k
ORDER=['Общие','Слова: Звери','Слова: Машины','Слова: Еда','Слова: Тело','Слова: Одежда','Слова: Посуда','Слова: Игрушки','Слова: Природа','Найди картинку','Кто это?','Сделай как я','Попугай учит','Договори','Что хочешь?','Считаем','Цветной коврик','Рисуем','Подуй']
NOTES={'q':'Вопрос: с интересом, интонация вверх в конце','wrong':'Спокойно, без огорчения, как подсказка','answer':'Ответ: уверенно и радостно',
 'sound':'Звукоподражание: игриво и выразительно, можно протяжно','word':'Название: чётко и радостно, как будто показываете картинку',
 'number':'Счёт: чётко и бодро, одинаково для всех чисел','cloze':'Начало фразы: голос «зависает», как будто ждёте, что ребёнок договорит. Без завершающей интонации',
 'praise':'Похвала: тепло и радостно, но не громко',
 'piece1':'Кусочек вопроса, начало выбора: интонация перечисления (вверх). Потом склеивается с «…или …?»',
 'piece2':'Кусочек вопроса, середина перечисления: ровно, как в списке','piece3':'Кусочек вопроса, конец: интонация вопроса вверх'}
def note(t,k):
    if k in NOTES: return NOTES[k]
    if t.endswith('?'): return NOTES['q']
    if t.startswith(('Молодец','Ура','Как красиво','Спасибо','Хорошо','Да!','Приехали','Приплыли','Свечк','Как крутится')): return NOTES['praise']
    if t.endswith('…'): return NOTES['cloze']
    if t.endswith('.'): return 'Спокойно и мягко'
    return 'Просьба или фраза: дружелюбно, медленно и чётко'
items=[]
for t,v in seen.items():
    prim='Общие' if 'Общие' in v['secs'] else v['secs'][0]
    items.append((ORDER.index(prim) if prim in ORDER else 99,prim,t,v['kind'],[s for s in v['secs'] if s!=prim]))
items.sort(key=lambda x:x[0])
wb=Workbook(); ws=wb.active; ws.title='Фразы'
F=Font(name='Arial',size=11); FB=Font(name='Arial',size=11,bold=True,color='FFFFFF'); HF=PatternFill('solid',start_color='3F5BA9')
BR=Border(bottom=Side(style='thin',color='D5DCE8'))
ws.append(['ID','Раздел игры','Фраза (что сказать)','Как произносить','Где ещё звучит','Файл','Готово','Комментарий'])
for c in ws[1]: c.font=FB; c.fill=HF; c.alignment=Alignment(vertical='center',wrap_text=True)
TINT=['FFF6D6','E8F4FF','E8F4FF','E8F4FF','E8F4FF','E8F4FF','E8F4FF','E8F4FF','E8F4FF','EFE9FF','FFEAF2','EAFBEF','FFF0E0','F2F2F2','E8F7FF','FFF9E0','FFE9E9','EEF9E6','F3EEFF']
SECFILL={s:PatternFill('solid',start_color=c) for s,c in zip(ORDER,TINT)}
for i,(o,prim,t,k,others) in enumerate(items,1):
    pid=f'p{i:04d}'; ws.append([pid,prim,t,note(t,k),', '.join(others),f'{pid}.mp3','',''])
    r=ws.max_row
    for c in ws[r]: c.font=F; c.alignment=Alignment(vertical='top',wrap_text=True); c.border=BR
    if prim in SECFILL: ws.cell(r,2).fill=SECFILL[prim]
    ws.cell(r,3).font=Font(name='Arial',size=12,bold=True)
for col,w in zip('ABCDEFGH',[8,18,42,46,24,12,9,30]): ws.column_dimensions[col].width=w
ws.freeze_panes='C2'; ws.auto_filter.ref=f"A1:H{ws.max_row}"
dv=DataValidation(type='list',formula1='"✓"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'G2:G{ws.max_row}')
last=ws.max_row
wi=wb.create_sheet('Как записывать')
lines=[('Озвучка игры «Покажи и скажи»',True),('',False),
('Для кого: малыши 3–5 лет, которые только начинают говорить. Голос должен быть тёплым, спокойным и улыбчивым, как у воспитателя, а не у диктора новостей.',False),('',False),
('Как говорить',True),
('• Медленнее обычного (примерно на 15–20%) и очень чётко, с ясными гласными.',False),
('• Ударения ставить естественно; буква «ё» в тексте всегда звучит как «ё».',False),
('• Похвала радостная, но не громкая и не визгливая: многие дети чувствительны к резким звукам.',False),
('• Вопросы — с живым интересом и подъёмом интонации в конце.',False),
('• Фразы с многоточием в конце («Раз, два…») — голос «зависает», чтобы ребёнок мог договорить сам.',False),
('• Кусочки вопросов «Хочешь мячик…», «…или машинку?» игра склеивает между собой, поэтому их важно записать ровно, без длинных пауз по краям.',False),
('• Колонка «Как произносить» — подсказка для каждой фразы.',False),('',False),
('Технические требования',True),
('• Один файл на одну строку, имя файла — из колонки «Файл» (например, p0001.mp3).',False),
('• Формат: MP3 192 kbps или WAV 44,1/48 кГц, моно.',False),
('• Тишина в начале и в конце — не больше 0,1 секунды.',False),
('• Одинаковая громкость у всех файлов (нормализация примерно до −16 LUFS, пики не выше −1 dB).',False),
('• Записывать в тихой комнате, без эха и фоновой музыки.',False),('',False),
('Как отмечать работу',True),
('• Заполняйте только колонки «Готово» и «Комментарий» на листе «Фразы».',False),
('• «Готово»: поставьте ✓ (можно выбрать из списка), когда файл записан. Например: строка p0001 → ✓.',False),
('• «Комментарий»: если фразу лучше сказать иначе или есть вопрос по ударению.',False),
('• Лист «Сводка» сам посчитает, сколько готово по каждому разделу.',False),('',False),
('Если озвучивает нейросеть (синтез речи)',True),
('• Выберите один голос для всех фраз и проверьте, что тариф разрешает использовать записи в коммерческом приложении.',False),
('• Колонку «Фраза» можно загрузить целиком, а «Как произносить» использовать как подсказку для стиля, если сервис это поддерживает.',False),
('• Обязательно прослушайте числа, звукоподражания («Мяу, мяу!», «Ко, ко, ко!») и фразы с «ё»: синтез чаще всего ошибается именно в них.',False)]
for t,b in lines:
    wi.append([t]); c=wi.cell(wi.max_row,1); c.font=Font(name='Arial',size=14 if (b and wi.max_row==1) else 11,bold=b); c.alignment=Alignment(wrap_text=True,vertical='top')
wi.column_dimensions['A'].width=110
wsum=wb.create_sheet('Сводка'); wsum.append(['Раздел','Фраз','Готово','Осталось'])
for c in wsum[1]: c.font=FB; c.fill=HF
for s in ORDER:
    r=wsum.max_row+1
    wsum.append([s,f'=COUNTIF(Фразы!$B$2:$B${last},A{r})',f'=COUNTIFS(Фразы!$B$2:$B${last},A{r},Фразы!$G$2:$G${last},"✓")',f'=B{r}-C{r}'])
r=wsum.max_row+1; wsum.append(['Всего',f'=SUM(B2:B{r-1})',f'=SUM(C2:C{r-1})',f'=B{r}-C{r}'])
for row in wsum.iter_rows(min_row=2):
    for c in row: c.font=F
for c in wsum[wsum.max_row]: c.font=Font(name='Arial',size=11,bold=True)
wsum.column_dimensions['A'].width=24
for col in 'BCD': wsum.column_dimensions[col].width=12
wb.move_sheet('Сводка',offset=-1); wb.move_sheet('Как записывать',offset=-2); wb.active=0
wb.save(dst); print(last-1,'phrases')
